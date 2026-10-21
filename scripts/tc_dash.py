"""多图拼版：dashboard（网格/自定义区域/多层共用横轴）与 facet（行×列分面网格）。
所有元素仍是可编辑 PPT 形状。拼版规范见 docs/STYLE_SPEC.md §10（[约定]）。
spec 结构见 docs/CATALOG.md「多图拼版」与 examples/dashboards/*.json。"""
import math
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
import tc_core
from tc_core import Slide, H, ink, set_run_font, TC

# 运营看板常用的蓝/橙/绿/灰配色：取色自用户提供的看板参考图（[约定]，仅在 spec 写 "palette": "ops_bright" 时使用）
DASH_PALETTES = {
    "ops_bright": ["#0070C0", "#EA6A06", "#4EA72E", "#BFBFBF", "#000000", "#9CAECC", "#4472C0", "#7F7F7F"],
}
NAMED = {"red": "#C00000", "green": "#3F9C1E", "gray": "#A6A6A6", "grey": "#A6A6A6", "blue": "#0070C0", "orange": "#EA6A06",
         "black": "#000000", "navy": "#1F3864", "lightgray": "#D9D9D9", "yellow": "#FFFF00"}
def C(c, default="#000000"):
    if c is None: return default
    return NAMED.get(c, c)
def darker(hexc, f=0.7):
    h = hexc.lstrip("#"); r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4))
    return "#%02X%02X%02X" % (int(r*f), int(g*f), int(b*f))
def inkhex(c): return "#FFFFFF" if str(ink(c)) == "FFFFFF" else "#000000"
def restrained(sl):
    return sl.theme.get("meta", {}).get("restrained", False)

def default_title_style(sl):
    return "plain" if restrained(sl) else "red_underline"

def text_col(c):
    """系列名/数值标签用的文字色：浅色系列（灰）自动加深，保证白底可读"""
    h = c.lstrip("#"); r, g, b = (int(h[i:i+2], 16) for i in (0, 2, 4))
    return darker(c, 0.62) if (0.299*r+0.587*g+0.114*b) > 170 else c

# ------------------------------------------------------------ 文字工具（多色分段）
def rich(sl, x, y, w, h, segs, size=9, align="l", anchor="m", bold=False, italic=False, color="#000000", wrap=False, fill=None):
    """segs: str 或 [{"text","color","bold","italic","fill"?}]；多行用 \\n。fill 给整个文本框底色（标签高亮块）"""
    if isinstance(segs, str): segs = [{"text": segs}]
    t = sl.s.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.05)), Inches(max(h, 0.05)))
    if fill: t.fill.solid(); t.fill.fore_color.rgb = H(fill)
    tf = t.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.03); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"m": MSO_ANCHOR.MIDDLE, "t": MSO_ANCHOR.TOP, "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]; p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for sg in segs:
        parts = str(sg.get("text", "")).split("\n")
        for k, part in enumerate(parts):
            if k: p = tf.add_paragraph(); p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            r = p.add_run(); r.text = part; set_run_font(r, sl.font); f = r.font; f.size = Pt(sg.get("size", size))
            f.bold = sg.get("bold", bold); f.italic = sg.get("italic", italic); f.color.rgb = H(C(sg.get("color"), color))
            if sg.get("underline"): f.underline = True
    return t
def est_w(s, size): return max(sum(1.0 if ord(ch) > 0x2E80 else 0.62 for ch in ln) for ln in str(s).split("\n")) * size / 72 + 0.16

def smooth(pts, k=8):
    """Catmull-Rom 采样成折线（PPT 自由曲线只支持直线段）"""
    if len(pts) < 3: return pts
    P = [pts[0]] + pts + [pts[-1]]; out = []
    for i in range(1, len(P)-2):
        p0, p1, p2, p3 = P[i-1], P[i], P[i+1], P[i+2]
        for j in range(k):
            t = j/k; t2, t3 = t*t, t*t*t
            out.append(tuple(0.5*((2*p1[a]) + (-p0[a]+p2[a])*t + (2*p0[a]-5*p1[a]+4*p2[a]-p3[a])*t2 + (-p0[a]+3*p1[a]-3*p2[a]+p3[a])*t3) for a in (0, 1)))
    out.append(pts[-1]); return out

def fmtv(v, d=None, suffix=""):
    if v is None: return ""
    if d is None: d = 0 if (abs(v) >= 100 or float(v).is_integer()) else 1
    return f"{v:,.{d}f}{suffix}"

# ------------------------------------------------------------ 面板（一个子图的坐标系）
class Panel:
    def __init__(self, sl, box, ch, pal, colors, n, pad=None):
        self.sl, self.ch, self.pal, self.colors = sl, ch, pal, colors
        x, y, w, h = box; p = dict(l=0.12, r=0.9, t=0.12, b=0.25); p.update(pad or {}); p.update(ch.get("pad", {}))
        if not ch.get("legend_right", True): p["r"] = min(p["r"], 0.12)
        if not ch.get("x_labels", True): p["b"] = min(p["b"], 0.06)
        self.box = box; self.px, self.pw = x+p["l"], w-p["l"]-p["r"]; self.py, self.ph = y+p["t"], h-p["t"]-p["b"]
        self.n = n; self.slot = self.pw/max(n, 1); self.fs = ch.get("font_size", 7.5); self.pts = {}
    def xs(self, i): return self.px + self.slot*(i+0.5)
    def set_y(self, lo, hi):
        if hi == lo: hi = lo + 1
        self.lo, self.hi = lo, hi
    def yy(self, v): return self.py + self.ph - (v-self.lo)/(self.hi-self.lo)*self.ph
    def col(self, name, k):
        return C(self.colors.get(name)) if name in self.colors else self.pal[k % len(self.pal)]
    def label(self, x, y, s, color="#000000", size=None, bold=False, italic=False, align="c", w=None):
        size = size or self.fs; w = w or est_w(s, size); hh = size/72*1.25*(str(s).count("\n")+1)
        x0 = x-w/2 if align == "c" else (x if align == "l" else x-w)
        rich(self.sl, x0, y-hh/2, w, hh, str(s), size=size, align=align, color=color, bold=bold, italic=italic)
        self.sl.mark(str(s), x, y, w, hh); return (x0, y-hh/2, w, hh)

def _right_names(P, items):
    """系列名放在右侧、用系列色；纵向自动错开避免重叠。items: [(y, name, color)]"""
    items = sorted(items); minh = P.fs/72*1.35; ys = []
    for y, nm, c in items:
        if ys and y - ys[-1] < minh: y = ys[-1] + minh
        ys.append(y); P.label(P.px+P.pw+0.06, y, nm, color=text_col(c), bold=True, align="l", size=P.ch.get("legend_size", P.fs+1.5), w=P.box[0]+P.box[2]-(P.px+P.pw)+0.4)

def _xlabels(P, cats):
    if not P.ch.get("x_labels", True): return
    every = P.ch.get("x_label_every", 1); sz = P.ch.get("x_label_size", P.fs-0.5)
    for i, c in enumerate(cats):
        if i % every == 0: P.label(P.xs(i), P.py+P.ph+0.12, c, size=sz, color="#404040", bold=P.ch.get("x_label_bold", True), w=max(P.slot*1.05, 0.3))
    if P.ch.get("baseline", True): P.sl.line(P.px, P.py+P.ph, P.px+P.pw, P.py+P.ph, w=0.5, color="#7F7F7F")

def area_stack(P, d):
    """堆叠面积：第 1 个系列在最下面；段内 '值 (占比%)'；顶部总计；右侧系列名"""
    cats, S = d["categories"], d["series"]; n = len(cats); pct = P.ch.get("percent", False)
    tot = [sum(s["values"][i] or 0 for s in S) for i in range(n)]
    P.set_y(0, 100 if pct else max(tot)*P.ch.get("headroom", 1.18))
    cum = [0.0]*n; right = []; sm = P.ch.get("smooth", False)
    for k, s in enumerate(S):
        v = [(x or 0)/(tot[i] or 1)*100 if pct else (x or 0) for i, x in enumerate(s["values"])]
        lo = [P.yy(c) for c in cum]; hi_v = [c+x for c, x in zip(cum, v)]; hi = [P.yy(c) for c in hi_v]
        xs = [P.xs(0)-P.slot*0.5] + [P.xs(i) for i in range(n)] + [P.xs(n-1)+P.slot*0.5]
        top = [(xs[0], hi[0])] + [(P.xs(i), hi[i]) for i in range(n)] + [(xs[-1], hi[-1])]
        bot = [(xs[0], lo[0])] + [(P.xs(i), lo[i]) for i in range(n)] + [(xs[-1], lo[-1])]
        if sm: top, bot = smooth(top), smooth(bot)
        c = P.col(s["name"], k); P.sl.poly(top + bot[::-1], fill=c, outline="#FFFFFF", lw=0.5)
        if P.ch.get("segment_labels", True):
            for i in range(n):
                hpx = lo[i]-hi[i]
                if hpx < P.ch.get("min_label_h", 0.16) or s["values"][i] is None: continue
                share = (s["values"][i] or 0)/(tot[i] or 1)*100
                txt = fmtv(s["values"][i], P.ch.get("decimals")) + (f"\n({share:.0f}%)" if P.ch.get("show_share", True) and not pct and hpx > 0.28 else "")
                if pct: txt = f"{share:.0f}%"
                P.label(P.xs(i), (lo[i]+hi[i])/2, txt, color=inkhex(c), size=P.fs-0.5, bold=True)
        right.append(((lo[-1]+hi[-1])/2, s["name"], c)); cum = hi_v
    if P.ch.get("totals", not pct):
        for i in range(n): P.label(P.xs(i), P.yy(tot[i])-0.11, fmtv(tot[i], P.ch.get("decimals")), bold=True, size=P.fs+1)
    if P.ch.get("legend_right", True): _right_names(P, right)
    _xlabels(P, cats)
    for i in range(n): P.pts[(None, i)] = (P.xs(i), P.yy(tot[i]))

def lines(P, d):
    """折线组：点上标数值（系列色）、右侧系列名、可平滑；label_pos: above|below|auto"""
    cats, S = d["categories"], d["series"]; n = len(cats); suf = P.ch.get("suffix", "")
    vals = [v for s in S for v in s["values"] if v is not None]
    lo, hi = P.ch.get("y_range", [min(vals), max(vals)]); pad = (hi-lo)*0.12 or 1
    P.set_y(lo-pad*P.ch.get("pad_lo", 1), hi+pad*P.ch.get("pad_hi", 1))
    right = []; sm = P.ch.get("smooth", False)
    for k, s in enumerate(S):
        c = P.col(s["name"], k); pts = [(P.xs(i), P.yy(v)) for i, v in enumerate(s["values"]) if v is not None]
        line_pts = smooth(pts) if sm else pts
        if len(line_pts) > 1: P.sl.poly(line_pts, outline=c, lw=s.get("width", P.ch.get("line_width", 1.75)), closed=False)
        pos = s.get("label_pos", P.ch.get("label_pos", "above"))
        for i, v in enumerate(s["values"]):
            if v is None: continue
            x, y = P.xs(i), P.yy(v); P.pts[(s["name"], i)] = (x, y)
            if P.ch.get("markers", False): P.sl.rect(x-0.03, y-0.03, 0.06, 0.06, c, shape=MSO_SHAPE.OVAL)
            if s.get("labels", P.ch.get("labels", True)):
                p = pos if pos != "auto" else ("above" if k % 2 == 0 else "below")
                P.label(x, y + (-1 if p == "above" else 1)*P.fs/72*0.95, fmtv(v, P.ch.get("decimals"), suf), color=text_col(c), size=P.fs, bold=True)
        if P.ch.get("legend_right", True) and pts: right.append((pts[-1][1], s["name"], c))
    if P.ch.get("legend_right", True): _right_names(P, right)
    _xlabels(P, cats)

def column_stack(P, d):
    """堆叠柱：段内标签（深底白字）、顶部总计；point_colors 可按点覆盖颜色（如平销/大促深浅不同）"""
    cats, S = d["categories"], d["series"]; n = len(cats)
    tot = [sum(s["values"][i] or 0 for s in S) for i in range(n)]; P.set_y(0, max(tot)*P.ch.get("headroom", 1.15))
    bw = P.slot*P.ch.get("bar_width", 0.62); right = []
    for i in range(n):
        y = 0
        for k, s in enumerate(S):
            v = s["values"][i] or 0; c = C((s.get("point_colors") or [None]*n)[i]) if (s.get("point_colors") or [None]*n)[i] else P.col(s["name"], k)
            P.sl.rect(P.xs(i)-bw/2, P.yy(y+v), bw, P.yy(y)-P.yy(y+v), c)
            if P.ch.get("segment_labels", True) and P.yy(y)-P.yy(y+v) > P.fs/72*1.2: P.label(P.xs(i), (P.yy(y)+P.yy(y+v))/2, fmtv(v, P.ch.get("decimals")), color=inkhex(c), size=P.fs, bold=True)
            if i == n-1: right.append(((P.yy(y)+P.yy(y+v))/2, s["name"], P.col(s["name"], k)))
            y += v
        if P.ch.get("totals", True):
            P.label(P.xs(i), P.yy(tot[i])-0.1, fmtv(tot[i], P.ch.get("decimals")), bold=True, size=P.fs+0.5)
        P.pts[(None, i)] = (P.xs(i), P.yy(tot[i]))
    if P.ch.get("legend_right", True): _right_names(P, [(y, nm, c) for y, nm, c in right])
    _xlabels(P, cats)

def bars_pair(P, d):
    """成对条形（分面单元）：当前值按规则红/绿，对比期灰；值写在条内左侧"""
    rows = d["rows"] if "rows" in d else [d]; rule = P.ch.get("rule", {}); thr = rule.get("threshold"); mx = P.ch.get("x_max", 100)
    n = len(rows); slot = P.ph/n; bh = slot*P.ch.get("bar_h", 0.28); xx = lambda v: P.px + min(v, mx)/mx*P.pw
    for i, r in enumerate(rows):
        if r is None: continue
        cy = P.py + slot*(i+0.5)
        v, cpv = r.get("value"), r.get("compare")
        if v is not None:
            good = P.sl.theme["good"] if restrained(P.sl) else C("green")
            bad = P.sl.theme["bad"] if restrained(P.sl) else C("red")
            if thr is None:
                c = P.sl.theme["pos"] if restrained(P.sl) else C(rule.get("lt", "red"))
            else:
                c = C(rule.get("ge"), good) if v >= thr else C(rule.get("lt"), bad)
            P.sl.rect(P.px, cy-bh, xx(v)-P.px, bh, c); P.label(P.px+0.04, cy-bh/2, fmtv(v, 1, "%"), color=inkhex(c), size=P.fs, align="l")
        if cpv is not None:
            g = C(rule.get("compare"), P.sl.theme["accent5"] if restrained(P.sl) else C("gray")); P.sl.rect(P.px, cy, xx(cpv)-P.px, bh, g); P.label(P.px+0.04, cy+bh/2, fmtv(cpv, 1, "%"), color=inkhex(g), size=P.fs, align="l")
    P.sl.line(P.px, P.py-0.05, P.px, P.py+P.ph, w=0.75, color="#404040")

KINDS = {"area_stack": area_stack, "lines": lines, "column_stack": column_stack, "bars_pair": bars_pair}

# ------------------------------------------------------------ 面板装饰：分段、圈、箭头、批注、通道、进度线
def decorate(P, ch):
    sl = P.sl; x0, y0, w0, h0 = P.box
    sty = dict(color=sl.theme["accent6"] if restrained(sl) else "#1F3864", italic=not restrained(sl), bold=True, size=P.fs+0.5); sty.update(ch.get("segment_style", {}))
    segs = ch.get("segments", [])
    for k, s in enumerate(segs):
        a, b = s["from"], s["to"]; xa, xb = P.xs(a)-P.slot/2, P.xs(b)+P.slot/2
        if s.get("label"): P.label((xa+xb)/2, y0+sty.get("dy", 0.04), s["label"], color=sty["color"], italic=sty["italic"], bold=sty["bold"], size=sty["size"])
        if k < len(segs)-1:
            sl.line(xb, y0+sty.get("line_top", 0.0), xb, P.py+P.ph+0.02, w=sty.get("line_w", 0.75), color=sty.get("line_color", "#808080"), dash=DASH.DASH)
    for a in ch.get("annotations", []):
        t = a["type"]; col = C(a.get("color"), sl.theme["callout"] if restrained(sl) else C("red"))
        def at(series, i, dy=0.0):
            if (series, i) in P.pts: x, y = P.pts[(series, i)]
            else: x, y = P.xs(i), P.py + P.ph/2
            return x, y + dy
        if t == "circle":
            pts = [at(a.get("series"), i) for i in a["points"]]; lab = a.get("on_label", True)
            dy = (-1 if a.get("label_pos", "above") == "above" else 1)*P.fs/72*0.95 if lab else 0
            xs = [p[0] for p in pts]; ys = [p[1]+dy for p in pts]
            w = max(xs)-min(xs)+a.get("w", 0.42); h = max(ys)-min(ys)+a.get("h", 0.2)
            sl.rect((max(xs)+min(xs))/2-w/2, (max(ys)+min(ys))/2-h/2, w, h, None, outline=col, lw=a.get("lw", 1.5), shape=MSO_SHAPE.OVAL)
        elif t == "arrow":
            fx, fy = a["from"]; tx, ty = a["to"]   # 面板内比例坐标 [0..1]
            sl.line(x0+fx*w0, y0+fy*h0, x0+tx*w0, y0+ty*h0, w=a.get("lw", 1.75), color=col, tail="triangle")
        elif t == "note":
            v = a.get("value")
            if a.get("color"):
                col = C(a["color"])
            elif restrained(sl):
                col = sl.theme["bad"] if v is not None and v < 0 else sl.theme["callout"]
            else:
                col = C("green") if (v or 0) > 0 else C("red")
            txt = a.get("text", "")
            if v is not None: txt += ("\n" if a.get("newline", True) else " ") + f"{v:+.1f}{a.get('unit', 'pp')}"
            if "point" in a: x, y = at(a.get("series"), a["point"]); x += a.get("dx", 0); y += a.get("dy", -0.3)
            else: x, y = x0+a["at"][0]*w0, y0+a["at"][1]*h0
            P.label(x, y, txt, color=col, bold=True, italic=a.get("italic", not restrained(sl)), size=a.get("size", P.fs+0.5))
        elif t == "channel":   # 趋势通道：对某系列做线性拟合，上下平移 offset 画两条虚线
            s = next(x for x in P.ch_data["series"] if x["name"] == a["series"]); v = [(i, y) for i, y in enumerate(s["values"]) if y is not None]
            m = len(v); mx = sum(i for i, _ in v)/m; my = sum(y for _, y in v)/m
            k = sum((i-mx)*(y-my) for i, y in v)/(sum((i-mx)**2 for i, _ in v) or 1); b = my-k*mx
            ia, ib = -0.5, P.n-0.5
            for off in (a.get("offset", 3), -a.get("offset", 3)):
                sl.line(P.xs(ia), P.yy(k*ia+b+off), P.xs(ib), P.yy(k*ib+b+off), w=0.75, color=C(a.get("color", "#808080")), dash=DASH.DASH)
        elif t == "text":
            P.label(x0+a["at"][0]*w0, y0+a["at"][1]*h0, a["text"], color=col if a.get("color") else "#000000", bold=a.get("bold", False), italic=a.get("italic", False), size=a.get("size", P.fs))

def panel_title(sl, x, y, w, t, style=None, size=10, align="c"):
    if not t: return
    style = style or default_title_style(sl)
    if isinstance(t, dict): style = t.get("style", style); t = t["text"]
    col = {"red_underline": "#C00000", "red": "#C00000", "plain": "#000000", "bold": "#000000"}.get(style, style)
    seg = [dict(text=t, color=col, bold=True, italic=style == "red_underline", underline=style == "red_underline")]
    rich(sl, x, y, w, size/72*1.6, seg, size=size, align=align)

def draw_panel(sl, box, p, pal, colors, pad=None):
    ch = p["chart"]; kind = ch["kind"]; d = ch["data"]
    title_h = max(0.28, p.get("title_size", 10)/72*1.65) if p.get("title") else 0.0
    if p.get("title"): panel_title(sl, box[0], box[1], box[2], p["title"], p.get("title_style"), p.get("title_size", 10), p.get("title_align", "l" if restrained(sl) else "c"))
    insight_h = p.get("insight_height", 0.75) if p.get("insight") else 0.0
    if insight_h and box[3]-title_h-insight_h < 0.9:
        raise ValueError("Panel insight leaves too little chart space; enlarge the panel or split the slide.")
    b = (box[0], box[1]+title_h, box[2], box[3]-title_h-insight_h)
    n = len(d.get("categories", d.get("rows", [])))
    P = Panel(sl, b, ch, pal, colors, n, pad); P.ch_data = d
    if ch.get("segments"): P.py += 0.22; P.ph -= 0.22
    KINDS[kind](P, d); decorate(P, ch)
    if insight_h:
        iy = box[1]+box[3]-insight_h+0.06
        sl.line(box[0], iy, box[0]+box[2], iy, w=0.6, color="#C9D0D5")
        rich(sl, box[0]+0.04, iy+0.08, box[2]-0.08, insight_h-0.18,
             p["insight"], size=p.get("insight_size", 12), color="#24465C", wrap=True, anchor="t")
    return P

# ------------------------------------------------------------ 页面
def page_frame(sl, sp):
    T = sp.get("title"); y = 0.18; size = sp.get("title_size", 20)
    segs = []
    if sp.get("title_tag"):
        tg = sp["title_tag"]; segs.append(dict(text=tg["text"], bold=True, color=tg.get("color", "#000000"), underline=tg.get("underline", True)))
        segs.append(dict(text=sp.get("title_sep", " | "), color="#000000", bold=True))
    segs += ([dict(text=T)] if isinstance(T, str) else (T or []))
    full = "".join(s["text"] for s in segs); tw = est_w(full, size)*0.95
    al = sp.get("title_align", "l" if restrained(sl) else "c"); X0 = (tc_core.SW-tw)/2 if al == "c" else 0.4
    if sp.get("title_tag") and sp["title_tag"].get("fill"):   # 标签高亮块（如黄底）
        tagw = est_w(sp["title_tag"]["text"], size)*1.0
        sl.rect(X0+0.02, y+0.04, tagw, size/72*1.35, C(sp["title_tag"]["fill"]))
    # 左对齐标题使用整行宽度，避免中文回退字体比估算更宽而意外折行。
    title_w = tc_core.SW-X0-0.4 if al == "l" else tw+0.3
    rich(sl, X0, y, title_w, size/72*1.6, segs, size=size, bold=sp.get("title_bold", True), align="l")
    if sp.get("note"): rich(sl, tc_core.SW-4.6, 0.18 + size/72*1.7, 4.4, 0.3, sp["note"], size=7, align="r", wrap=True, color="#7F7F7F")
    fy = tc_core.SH-0.12; lines_ = sp.get("footnotes", [])
    src = sp.get("source")
    allf = ([[dict(text="Source: "), *([dict(text=src)] if isinstance(src, str) else src)]] if src else [[dict(text="Source: [missing – add source]", color="#C00000")]]) + [l if isinstance(l, list) else [dict(text=l)] for l in lines_]
    foot_size = sp.get("footnote_size", 9 if sp.get("exhibit") else 7)
    foot_h = max(0.16, foot_size/72*1.6)
    for k, l in enumerate(reversed(allf)):
        rich(sl, 0.35, fy-foot_h*(k+1), 12.5, foot_h, l, size=foot_size, color="#404040")
    top = 0.18 + size/72*1.7 + 0.05 + (0.32 if sp.get("note") else 0)
    if sp.get("exhibit"):
        sl.line(0.35, top+0.015, tc_core.SW-0.35, top+0.015, w=0.9, color=sl.theme["pos"])
        top += 0.20
    return top, fy-foot_h*len(allf)-0.05

def render_dashboard(prs, sp):
    sl = Slide(prs, sp); pal = DASH_PALETTES.get(sp.get("palette"), None) or (sp.get("palette") if isinstance(sp.get("palette"), list) else sl.pal)
    colors = sp.get("series_colors", {}); top, bot = page_frame(sl, sp)
    X, W = 0.3, tc_core.SW-0.6; Y, Hh = top, bot-top
    g = sp.get("grid", {"rows": 1, "cols": len(sp["panels"])}); gap = g.get("gap", 0.15)
    cw = (W-gap*(g["cols"]-1))/g["cols"]; rh = (Hh-gap*(g["rows"]-1))/g["rows"]
    pads = {}
    for p in sp["panels"]:
        if "box" in p: fx, fy, fw, fh = p["box"]; box = (X+fx*W, Y+fy*Hh, fw*W, fh*Hh)
        else:
            r, c = p.get("cell", [0, 0]); rs, cs = p.get("span", [1, 1])
            box = (X+c*(cw+gap), Y+r*(rh+gap), cs*cw+(cs-1)*gap, rs*rh+(rs-1)*gap)
        draw_panel(sl, box, p, pal, colors, sp.get("pad"))
    for s in sp.get("separators", []):   # 页面级分隔虚线：{"x":比例} 或 {"y":比例}
        if "x" in s: sl.line(X+s["x"]*W, Y, X+s["x"]*W, bot, w=0.75, color="#A6A6A6", dash=DASH.DASH)
        else: sl.line(X, Y+s["y"]*Hh, X+W, Y+s["y"]*Hh, w=0.75, color="#A6A6A6", dash=DASH.DASH)
    return sl

def render_facet(prs, sp):
    """分面网格：rows=指标（各自图型），cols=分组；共享 y（按行 shared_y="row"）或独立（"none"）"""
    sl = Slide(prs, sp); pal = DASH_PALETTES.get(sp.get("palette"), None) or (sp.get("palette") if isinstance(sp.get("palette"), list) else sl.pal)
    colors = sp.get("series_colors", {}); top, bot = page_frame(sl, sp)
    rows, cols = sp["rows"], sp["cols"]; nr, nc = len(rows), len(cols)
    rhw = sp.get("row_header_w", 1.3); lgw = sp.get("legend_w", 0.9); colh = sp.get("col_header_h", 0.35)
    X = 0.3 + rhw; W = tc_core.SW-0.3-X-lgw; Y = top+colh; Hh = bot-Y-sp.get("bottom_h", 0.0)
    cw, rh = W/nc, Hh/nr; chs = sp.get("col_header_style", default_title_style(sl)); rst = sp.get("row_header_style", "plain" if restrained(sl) else "box")
    for j, c in enumerate(cols):
        panel_title(sl, X+j*cw, top+0.02, cw, c, chs, sp.get("col_header_size", 11), sp.get("col_header_align", "c"))
    for i, r in enumerate(rows):
        y = Y+i*rh
        if rst == "box":
            sl.rect(0.35, y+rh*0.22, rhw-0.15, rh*0.56, "#FFFFFF", outline="#7F7F7F", lw=0.75)
            rich(sl, 0.35, y+rh*0.22, rhw-0.15, rh*0.56, r["label"], size=sp.get("row_header_size", 10), bold=True, italic=True, align="c", wrap=True)
        else:
            rich(sl, 0.3, y, rhw-0.1, rh, r["label"], size=sp.get("row_header_size", 10), bold=rst != "plain", align="r" if rst == "right" else "l", wrap=True)
        ch = dict(r.get("chart", {})); kind = ch.get("kind")
        cells = [sp["cells"][i][j] for j in range(nc)]
        shared = ch.get("shared_y", sp.get("shared_y", "row"))
        if shared == "row" and kind == "lines":
            vv = [v for d in cells if d for s in d["series"] for v in s["values"] if v is not None]; ch.setdefault("y_range", [min(vv), max(vv)])
        if kind == "bars_pair": ch.setdefault("x_max", 100)
        for j, d in enumerate(cells):
            if not d: continue
            cch = dict(ch); cch["data"] = d; cch["legend_right"] = (j == nc-1) and ch.get("legend_right", True)
            cch["x_labels"] = ch.get("x_labels", i == nr-1)
            pad = dict(l=0.08, r=0.08 if j < nc-1 else 0.08, t=0.08, b=0.22 if cch["x_labels"] else 0.06)
            if j == nc-1 and cch["legend_right"]: pad["r"] = 0.08
            box = (X+j*cw, y, cw, rh)
            P = Panel(sl, box, cch, pal, colors, len(d.get("categories", d.get("rows", [0]))), pad); P.ch_data = d
            if kind == "area_stack" and shared == "row" and not ch.get("percent"):
                mx = max(sum(s["values"][k] or 0 for s in dd["series"]) for dd in cells if dd for k in range(len(dd["categories"])))
                orig = P.set_y; P.set_y = lambda lo, hi, mx=mx: orig(0, mx*ch.get("headroom", 1.18))
            KINDS[kind](P, d)   # 最后一列的系列名画在 px+pw 右侧，即网格外的 legend 区
            decorate(P, cch)
    th_rows = [r["chart"] for r in rows if r.get("chart", {}).get("kind") == "bars_pair" and r["chart"].get("rule", {}).get("threshold") is not None]
    if th_rows:   # 时间进度虚线：每列一条贯穿全部行，底部标 ▲ 数值
        t = th_rows[0]["rule"]["threshold"]; mx = th_rows[0].get("x_max", 100)
        for j in range(nc):
            xt = X+j*cw+0.08 + t/mx*(cw-0.16)
            sl.line(xt, Y, xt, Y+Hh, w=0.75, color="#404040", dash=DASH.DASH)
            rich(sl, xt-0.45, Y+Hh+0.02, 0.9, 0.2, [dict(text="▲ " + fmtv(t, 1, "%"))], size=8, align="c", bold=True)
    sep = sp.get("separators", "dash")
    if sep:
        for i in range(1, nr): sl.line(0.3, Y+i*rh, X+W+lgw, Y+i*rh, w=0.75, color="#7F7F7F", dash=DASH.DASH)
        if sp.get("col_separators", True):
            for j in range(1, nc): sl.line(X+j*cw, top+0.05, X+j*cw, Y+Hh, w=0.5, color="#BFBFBF", dash=DASH.ROUND_DOT)
    return sl

def render(prs, spec):
    return render_facet(prs, spec) if spec["type"] == "facet" else render_dashboard(prs, spec)
