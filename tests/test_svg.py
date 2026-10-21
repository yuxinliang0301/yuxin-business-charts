"""HTML 矢量后端的覆盖、几何和只读回归测试。"""
import json
from pathlib import Path
import re
import sys
import unittest

from lxml import etree
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import tc_core
import tc_more  # 注册完整图型。
import tc_dash
import tc_svg

NS = {"s": tc_svg.SVG_NS, "h": tc_svg.XHTML_NS}


def parse(prs):
    return etree.fromstring(tc_svg.presentation_to_svg(prs)[0].encode())


def slide():
    prs = tc_core.new_prs()
    return prs, tc_core.Slide(prs, {})


class SVGExportTests(unittest.TestCase):
    def test_every_example_is_vector_editable_and_read_only(self):
        prs = tc_core.new_prs()
        examples = []
        for path in sorted((ROOT / "examples").rglob("*.json")):
            spec = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(spec, dict) and ("type" in spec or "slides" in spec):
                for item in spec.get("slides", [spec]):
                    tc_core.render(prs, item)
                    examples.append(path.name)
        # 覆盖完整示例集，防止把不认识的图型静默跳过。
        self.assertGreaterEqual(len(examples), 72)
        before = {str(p.partname): p.blob for p in prs.part.package.iter_parts()}
        svgs = tc_svg.presentation_to_svg(prs)
        after = {str(p.partname): p.blob for p in prs.part.package.iter_parts()}
        self.assertEqual(before, after, "HTML 导出不应修改任何 PPTX part")
        self.assertEqual(len(svgs), len(examples))
        ids = set()
        for index, value in enumerate(svgs):
            root = etree.fromstring(value.encode())
            self.assertFalse(root.xpath(".//s:image | .//s:script", namespaces=NS))
            groups = root.xpath("./s:g[@data-shape-id]", namespaces=NS)
            self.assertEqual(len(groups), len(prs.slides[index].shapes))
            for div in root.xpath(".//h:div[@class='tc-editable']", namespaces=NS):
                ident = div.get("data-edit-id")
                self.assertNotIn(ident, ids)
                ids.add(ident)
                self.assertEqual(div.get("contenteditable"), "false")
        self.assertGreater(len(ids), 3000)

    def test_primitive_geometry_and_point_units(self):
        prs, sl = slide()
        sl.rect(1, 2, 3, 1, "#123456", outline="#654321", lw=2)
        sl.rect(1, 1, 2, 2, "#FFFFFF", shape=MSO_SHAPE.OVAL)
        sl.poly([(1, 1), (2, 3), (4, 1)], fill="#ABCDEF")
        root = parse(prs)
        self.assertEqual(root.get("viewBox"), "0 0 959.975984 540")
        groups = root.xpath("./s:g", namespaces=NS)
        self.assertEqual(groups[0].get("transform"), "translate(72 144)")
        rect = groups[0].find(".//s:rect", NS)
        self.assertEqual((rect.get("width"), rect.get("height")), ("216", "72"))
        self.assertEqual(rect.getparent().get("fill"), "#123456")
        self.assertEqual(rect.getparent().get("stroke-width"), "2")
        ellipse = groups[1].find(".//s:ellipse", NS)
        self.assertEqual(ellipse.get("rx"), "72")
        path = groups[2].find(".//s:path", NS)
        self.assertEqual(path.get("d"), "M 0 0 L 72 144 L 216 0 Z")

    def test_sector_direction_ring_thickness_and_full_circle(self):
        prs, sl = slide()
        pie = sl.rect(0, 0, 2, 2, "#000000", shape=MSO_SHAPE.PIE)
        pie.adjustments[0], pie.adjustments[1] = 270 * .6, 90 * .6
        ring = sl.rect(3, 0, 2, 2, "#000000", shape=MSO_SHAPE.BLOCK_ARC)
        ring.adjustments[0], ring.adjustments[1], ring.adjustments[2] = 270 * .6, 90 * .6, .18
        circle = sl.rect(6, 0, 2, 2, "#000000", shape=MSO_SHAPE.PIE)
        circle.adjustments[0], circle.adjustments[1] = 0, 0
        paths = parse(prs).xpath("./s:g/s:g/s:path/@d", namespaces=NS)
        # 十二点开始顺时针到六点，填右半圆。
        self.assertEqual(paths[0], "M 72 0 A 72 72 0 0 1 72 144 L 72 72 Z")
        # 厚度为 144 * .18 = 25.92 point，内半径为 46.08。
        self.assertEqual(paths[1], "M 72 0 A 72 72 0 0 1 72 144 L 72 118.08 A 46.08 46.08 0 0 0 72 25.92 Z")
        self.assertEqual(paths[2].count("A 72 72"), 2)

    def test_reversed_connectors_keep_arrow_ends_and_dashes(self):
        prs, sl = slide()
        sl.line(3, 2, 1, 1, w=2, color="#123456", dash=MSO_LINE_DASH_STYLE.DASH, head="arrow", tail="triangle")
        root = parse(prs)
        group = root.find("./s:g", NS)
        self.assertEqual(group.get("transform"), "translate(72 72) translate(144 72) scale(-1 -1)")
        line = group.find("./s:g", NS)
        self.assertEqual(line.get("marker-start"), "url(#s0-shape2-headEnd)")
        self.assertEqual(line.get("marker-end"), "url(#s0-shape2-tailEnd)")
        self.assertEqual(line.get("stroke-dasharray"), "8 6")
        markers = root.xpath(".//s:marker", namespaces=NS)
        self.assertTrue(all(m.get("orient") == "auto-start-reverse" for m in markers))

    def test_rich_text_safe_escape_format_and_edit_inheritance(self):
        prs, sl = slide()
        label = '<script>alert("x")</script> & 中文\n第二行'
        tc_dash.rich(sl, 1, 1, 3, 1, [{"text": label, "color": "#123456", "bold": True, "italic": True, "underline": True}, {"text": "尾部", "size": 17}], size=14, align="c", anchor="b", wrap=True, fill="#FFFF00")
        root = parse(prs)
        self.assertFalse(root.xpath(".//s:script | .//h:script", namespaces=NS))
        div = root.find(".//h:div", NS)
        self.assertIn("text-align:center", div.get("style"))
        self.assertIn("line-height:1.15", div.get("style"))
        self.assertIn("justify-content:flex-end", div.get("style"))
        self.assertIn("white-space:pre-wrap", div.get("style"))
        self.assertIn("padding:0px 2.16px 0px 2.16px", div.get("style"))
        self.assertIn("font-size:14px", div.get("style"))
        runs = div.findall(".//h:span", NS)
        self.assertEqual(runs[0].text, label.split("\n")[0])
        self.assertIn("font-weight:700", runs[0].get("style"))
        self.assertIn("font-style:italic", runs[0].get("style"))
        self.assertIn("text-decoration:underline", runs[0].get("style"))
        self.assertIn("font-size:17px", runs[-1].get("style"))
        self.assertEqual(len(div.findall("h:p", NS)), 2)
        self.assertEqual(root.find("./s:g/s:g", NS).get("fill"), "#FFFF00")

    def test_unknown_shape_fails_instead_of_disappearing(self):
        prs, sl = slide()
        sl.s.shapes.add_shape(MSO_SHAPE.HEART, Inches(1), Inches(1), Inches(1), Inches(1))
        with self.assertRaisesRegex(tc_svg.UnsupportedShapeError, "heart"):
            tc_svg.presentation_to_svg(prs)


if __name__ == "__main__":
    unittest.main()
